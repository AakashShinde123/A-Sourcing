#!/usr/bin/env python3
"""Generate sample asset-register workbooks for EasySourcing import testing.

Files (in download/sample-data/):
  - Meridian-Asset-Register-SAMPLE.xlsx        12-row register for Meridian Manufacturing (MRD)
  - Meridian-Register-ROUND2-SAMPLE.xlsx       dedup demo: 3 repeats of file 1 + 3 new rows
  - (CSV mirrors of both for direct drop)

Column names ride the parser's alias table (Asset ID / Description / Category /
Make / Model / Serial No / Barcode / Location / Custodian). Locations are real
Meridian location names so the import's fuzzy matcher links them.
"""
import os, sys

XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from templates.base import (
    setup_sheet, style_header_row, style_data_row, font_body, font_caption,
    align_text, auto_fit_columns, NEUTRAL_600,
)

OUT = "/home/z/my-project/download/sample-data"
os.makedirs(OUT, exist_ok=True)

HEADERS = ["Asset ID", "Description", "Category", "Make", "Model", "Serial No", "Barcode", "Location", "Custodian"]

# 12 fresh assets for Meridian Manufacturing (MRD). None exist in the platform yet.
ROWS_1 = [
    ["MRD-FA-2201", "Vertical Machining Center VMC-850", "Production Machinery", "Bharat Fritz Werner", "VMC-850", "BFW-88213",  "BC-MRD-1201", "CNC Hall 1",            "R. Tambat"],
    ["MRD-FA-2202", "CNC Turning Center ST-20",         "Production Machinery", "ACE Designers",       "LT-20",    "ACE-44190", "BC-MRD-1202", "CNC Hall 2",            "S. Pawar"],
    ["MRD-FA-2203", "Hydraulic Pallet Truck 2.5T",      "Material Handling",    "Godrej",              "HPD-25",   "GD-33077",  "BC-MRD-1203", "FG Bay 1",              "D. Joshi"],
    ["MRD-FA-2204", "Server Rack 42U with PDU",         "IT Infrastructure",    "APC",                 "SR42U",    "AP-77231",  "BC-MRD-1204", "Server Room",           "S. Iyer"],
    ["MRD-FA-2205", "Air Compressor 15 kW",             "Utilities",            "Atlas Copco",         "GA-15",    "AC-20988",  "BC-MRD-1205", "Workshop",              "M. Bhosale"],
    ["MRD-FA-2206", "Diesel Generator 125 kVA",         "Utilities",            "Kirloskar",           "KG1-125AS", "KO-51120", "BC-MRD-1206", "Maintenance Department","V. Kale"],
    ["MRD-FA-2207", "Assembly Conveyor Section C",      "Production Machinery", "Flexlink",            "FL-C12",   "FX-12045",  "BC-MRD-1207", "Assembly Line 1",       "P. Shinde"],
    ["MRD-FA-2208", "Automatic Band Sealer",            "Packaging",            "Sealpac",             "SP-900",   "SP-66334",  "BC-MRD-1208", "Finished Goods Store",  "A. More"],
    ["MRD-FA-2209", "3-Axis Coordinate Measuring Machine","Quality & Inspection","Carl Zeiss",         "Contura",  "CZ-30071",  "BC-MRD-1209", "Workshop",              "K. Menon"],
    ["MRD-FA-2210", "Warehouse Storage Rack Bay 12",    "Storage & Racking",    "Godrej",              "GR-B12",   None,        "BC-MRD-1210", "RM Bay",                "N. Gaikwad"],
    ["MRD-FA-2211", "Fire Extinguisher Cabinet B2",     "Safety Equipment",     "Safequip",            "SQ-B2",    None,        None,          "Cafeteria Block C",     "Facilities Team"],
    ["MRD-FA-2212", "Design Workstation Dell 7550",     "IT Infrastructure",    "Dell",                "Precision 7550", "DL-91256", "BC-MRD-1212", "Design Studio",     "A. Rahane"],
]

# Round 2: first 3 repeat Round 1 (dedup demo), last 3 are genuinely new.
ROWS_2 = [
    ["MRD-FA-2201", "Vertical Machining Center VMC-850", "Production Machinery", "Bharat Fritz Werner", "VMC-850",  "BFW-88213", "BC-MRD-1201", "CNC Hall 1",   "R. Tambat"],
    ["MRD-FA-2202", "CNC Turning Center ST-20",          "Production Machinery", "ACE Designers",       "LT-20",    "ACE-44190", "BC-MRD-1202", "CNC Hall 2",   "S. Pawar"],
    ["MRD-FA-2204", "Server Rack 42U with PDU",          "IT Infrastructure",    "APC",                 "SR42U",    "AP-77231",  "BC-MRD-1204", "Server Room",  "S. Iyer"],
    ["MRD-FA-2213", "Battery Operated Pallet Stacker",   "Material Handling",    "Godrej",              "PS-1500",  "GD-41002",  "BC-MRD-1213", "FG Bay 1",     "H. Patil"],
    ["MRD-FA-2214", "Precision Surface Grinder",         "Production Machinery", "Paragon Machinery",   "PSG-52",   "PG-20991",  "BC-MRD-1214", "CNC Hall 1",   "R. Tambat"],
    ["MRD-FA-2215", "Network Switch Rack 24-Port",       "IT Infrastructure",    "Cisco",               "Catalyst 9200", "CS-88410", "BC-MRD-1215", "Server Room", "S. Iyer"],
]

NOTES = [
    "Sample register for import testing · Meridian Manufacturing Pvt Ltd (client code MRD)",
    "Round 2 file intentionally repeats 3 rows from the Round 1 file — the platform must report them as duplicates and skip them.",
    "Rows without Serial No / Barcode show that optional columns may be left blank.",
    "MRD-FA-2211 points at 'Cafeteria Block C' which is NOT a Meridian location — the platform should import the asset but leave its location unlinked.",
]


def build(path_base: str, rows: list, notes: list) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Asset Register"

    setup_sheet(ws, title="Meridian Manufacturing — Fixed Asset Register", last_col=1 + len(HEADERS))

    header_row = 4
    for col, h in enumerate(HEADERS, 2):
        ws.cell(row=header_row, column=col, value=h)
    style_header_row(ws, header_row, 2, 1 + len(HEADERS))

    for i, row in enumerate(rows):
        r = header_row + 1 + i
        for col, v in enumerate(row, 2):
            c = ws.cell(row=r, column=col, value=v if v is not None else "")
            c.alignment = align_text()
        style_data_row(ws, r, 2, 1 + len(HEADERS), i)
        # restore body font lost to style_data_row defaults on empty cells
        for col in range(2, 2 + len(HEADERS)):
            if ws.cell(row=r, column=col).value == "":
                ws.cell(row=r, column=col).font = font_body()

    note_row = header_row + len(rows) + 3
    for j, note in enumerate(notes):
        cell = ws.cell(row=note_row + j, column=2, value=("Notes: " if j == 0 else "") + note)
        cell.font = font_caption()

    auto_fit_columns(ws, header_row=header_row, data_start_row=header_row + 1)
    ws.freeze_panes = "B5"

    wb.properties.creator = "Z.ai"
    xlsx_path = os.path.join(OUT, path_base + ".xlsx")
    wb.save(xlsx_path)

    # CSV mirror — drops straight into the Import dialog with zero conversion
    csv_path = os.path.join(OUT, path_base + ".csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        f.write(",".join(HEADERS) + "\n")
        for row in rows:
            f.write(",".join("" if v is None else str(v) for v in row) + "\n")
    print("wrote", xlsx_path)
    print("wrote", csv_path)


build("Meridian-Asset-Register-SAMPLE", ROWS_1, NOTES)
build("Meridian-Register-ROUND2-SAMPLE", ROWS_2, NOTES[:1] + [NOTES[1]])
print("done")
